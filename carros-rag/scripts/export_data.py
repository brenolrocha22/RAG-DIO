from __future__ import annotations

import argparse
import csv
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
import unicodedata

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


REQUIRED_COLUMNS = {
    "id",
    "marca",
    "modelo",
    "preco",
    "pais",
    "regiao",
    "nome",
    "data",
}


def normalize_header(value: Any) -> str:
    normalized = unicodedata.normalize("NFKD", str(value or ""))
    return "".join(char for char in normalized.casefold() if char.isalnum())


def cell_text(value: Any, column: str, workbook_epoch: datetime) -> str:
    if value is None:
        return ""
    if column == "data":
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        if isinstance(value, (int, float)):
            return from_excel(value, workbook_epoch).date().isoformat()
        try:
            return date.fromisoformat(str(value).strip()).isoformat()
        except ValueError as error:
            raise ValueError(f"Data inválida na planilha: {value!r}") from error
    return str(value).strip()


def write_csv(path: Path, fieldnames: list[str], rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as output:
        writer = csv.DictWriter(output, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def export_workbook(workbook_path: Path, output_dir: Path) -> tuple[int, int, int]:
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Planilha não encontrada: {workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if "base" not in workbook.sheetnames:
        raise ValueError("A planilha precisa conter uma aba chamada 'base'.")

    sheet = workbook["base"]
    iterator = sheet.iter_rows(values_only=True)
    try:
        header_row = next(iterator)
    except StopIteration as error:
        raise ValueError("A aba 'base' está vazia.") from error

    column_positions = {
        normalize_header(header): position
        for position, header in enumerate(header_row)
        if header is not None
    }
    missing_columns = REQUIRED_COLUMNS - column_positions.keys()
    if missing_columns:
        raise ValueError(
            "Colunas obrigatórias ausentes: " + ", ".join(sorted(missing_columns))
        )

    source_rows: list[dict[str, str]] = []
    for row_number, values in enumerate(iterator, start=2):
        if not any(value is not None for value in values):
            continue
        record = {
            column: cell_text(
                values[column_positions[column]]
                if column_positions[column] < len(values)
                else None,
                column,
                workbook.epoch,
            )
            for column in REQUIRED_COLUMNS
        }
        missing_values = [key for key, value in record.items() if not value]
        if missing_values:
            raise ValueError(
                f"Linha {row_number} sem valores obrigatórios: "
                + ", ".join(sorted(missing_values))
            )
        try:
            record["preco"] = format(Decimal(record["preco"]).quantize(Decimal("0.01")), "f")
        except InvalidOperation as error:
            raise ValueError(
                f"Preço inválido na linha {row_number}: {record['preco']!r}"
            ) from error
        source_rows.append(record)

    if not source_rows:
        raise ValueError("A aba 'base' não contém registros de vendas.")

    vehicles = sorted(
        {(row["marca"], row["modelo"]) for row in source_rows},
        key=lambda item: (item[0].casefold(), item[1].casefold()),
    )
    vehicle_ids = {
        vehicle: f"VEI-{position:04d}"
        for position, vehicle in enumerate(vehicles, start=1)
    }
    vehicle_rows = [
        {
            "vehicle_id": vehicle_ids[vehicle],
            "marca": vehicle[0],
            "modelo": vehicle[1],
        }
        for vehicle in vehicles
    ]

    customer_keys = sorted(
        {
            (row["nome"], row["pais"], row["regiao"])
            for row in source_rows
        },
        key=lambda item: tuple(value.casefold() for value in item),
    )
    customer_ids = {
        customer: f"CLI-{position:04d}"
        for position, customer in enumerate(customer_keys, start=1)
    }
    customer_rows = [
        {
            "customer_id": customer_ids[customer],
            "pais": customer[1],
            "regiao": customer[2],
        }
        for customer in customer_keys
    ]

    sales_rows = [
        {
            "sale_id": f"VEN-{position:06d}",
            "vehicle_id": vehicle_ids[(row["marca"], row["modelo"])],
            "customer_id": customer_ids[(row["nome"], row["pais"], row["regiao"])],
            "data_venda": row["data"],
            "preco": row["preco"],
        }
        for position, row in enumerate(source_rows, start=1)
    ]

    write_csv(
        output_dir / "vendas.csv",
        ["sale_id", "vehicle_id", "customer_id", "data_venda", "preco"],
        sales_rows,
    )
    write_csv(
        output_dir / "veiculos.csv",
        ["vehicle_id", "marca", "modelo"],
        vehicle_rows,
    )
    write_csv(
        output_dir / "clientes.csv",
        ["customer_id", "pais", "regiao"],
        customer_rows,
    )
    workbook.close()
    return len(sales_rows), len(vehicle_rows), len(customer_rows)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Gera os três CSVs anonimizados do projeto de RAG."
    )
    parser.add_argument("workbook", type=Path, help="Caminho da planilha .xlsx")
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path(__file__).resolve().parents[1] / "data",
        help="Pasta de destino dos CSVs (padrão: data/ do projeto)",
    )
    args = parser.parse_args()
    sales_count, vehicle_count, customer_count = export_workbook(
        args.workbook, args.output_dir
    )
    print(
        f"CSV criados: {sales_count} vendas, {vehicle_count} veículos e "
        f"{customer_count} clientes anonimizados em {args.output_dir.resolve()}"
    )


if __name__ == "__main__":
    main()

